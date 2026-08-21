"""Operational features: dry run, per-command retry, progress hook, manifest,
and the configurable knobs that used to be constants."""

import json
import shutil
from pathlib import Path

import pytest

from switch_migrator import manifest as manifest_mod
from switch_migrator.cli import build_arg_parser, main, preview_commands
from switch_migrator.collectors.switch import collect_switch
from switch_migrator.config import Config, SshSettings, SwitchTarget, load_config
from switch_migrator.connection import CommandError, DryRunRunner, OfflineRunner
from switch_migrator.models import FabricState, Platform
from switch_migrator.report.progress import NullProgress, make_progress

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture
def cfg() -> Config:
    return Config(dvr_controllers=[], core_switch_patterns=["dvr-*"],
                  isid_offsets=[10000], isid_explicit={}, excluded_vlans={1},
                  ssh=SshSettings())


# --------------------------- dry run ---------------------------------------

def test_dry_run_lists_commands_without_connecting(cfg):
    args = build_arg_parser().parse_args(["--migration-sheets"])
    preview = preview_commands(
        [SwitchTarget("gx-01", "10.0.0.1", Platform.VOSS),
         SwitchTarget("ers-01", "10.0.0.2", Platform.ERS)], cfg, args)
    voss, ers = preview["gx-01"], preview["ers-01"]
    # session setup first - 'enable' and EVERY paging spelling (fallbacks
    # included; the list is the full honest answer) - then the real reads
    assert voss[:3] == ["enable", "terminal more disable", "term more dis"]
    assert "show mlt" in voss and "show vlan i-sid" in voss
    assert ers[:2] == ["enable", "terminal length 0"]
    assert "show mac-address-table" in ers          # ERS FDB command
    assert "show interfaces gigabitEthernet fdb-entry" in voss   # VOSS FDB command
    # read-only is the claim the preview exists to support: only reads, the
    # privilege mode switch and terminal-paging settings
    for commands in preview.values():
        for command in commands:
            assert command == "enable" or \
                command.startswith(("show ", "terminal ", "term ")), command


def test_dry_run_covers_both_port_state_variants(cfg):
    """Empty answers walk every fallback branch, so the preview is the full
    set of commands that could be sent - not just one release's subset."""
    args = build_arg_parser().parse_args([])
    preview = preview_commands(
        [SwitchTarget("gx-01", "gx-01", Platform.VOSS)], cfg, args)
    assert "show interfaces gigabitEthernet state" in preview["gx-01"]
    assert "show interfaces gigabitEthernet interface" in preview["gx-01"]


def test_dry_run_config_pull_is_opt_in(cfg):
    plain = build_arg_parser().parse_args([])
    with_config = build_arg_parser().parse_args(["--extract-config"])
    target = [SwitchTarget("gx-01", "gx-01", Platform.VOSS)]
    assert "show running-config" not in preview_commands(target, cfg, plain)["gx-01"]
    assert "show running-config" in preview_commands(target, cfg, with_config)["gx-01"]


def test_dry_run_end_to_end_exits_clean(tmp_path, capsys):
    cfg_path = tmp_path / "config.yaml"
    cfg_path.write_text("dvr_controllers:\n  - {name: dvr-01, host: 10.0.0.1}\n"
                        "isid_conventions:\n  offsets: [10000]\n")
    code = main(["-c", str(cfg_path), "-o", str(tmp_path / "out"),
                 "-s", "gx-01:voss", "--dry-run"])
    assert code == 0
    out = capsys.readouterr().err
    assert "nothing was connected to" in out
    assert "show mlt" in out
    assert "dvr-01" in out          # the controller commands are previewed too


# --------------------------- per-command retry -----------------------------

class _FlakyConn:
    """Fails the first N send_command calls with a transport error."""

    def __init__(self, failures: int, output: str = "ok"):
        self.failures = failures
        self.output = output
        self.calls = 0

    def send_command(self, command, read_timeout=None):
        self.calls += 1
        if self.calls <= self.failures:
            raise TimeoutError("Pattern not detected: '(?:#|>)'")
        return self.output

    def write_channel(self, data):
        pass

    def clear_buffer(self):
        pass


def _runner(conn, **ssh_over):
    from switch_migrator import connection

    runner = object.__new__(connection.SshRunner)
    runner.name, runner.host = "sw", "sw"
    runner.platform = Platform.VOSS
    runner.ssh = SshSettings(**ssh_over)
    runner.raw_dir = None
    runner.setup_warnings = []
    runner.command_log = []
    runner.on_command = None
    runner._transport_failures = 0
    runner._dead = False
    runner._conn = conn
    return runner


def test_one_dropped_read_is_retried_not_lost(monkeypatch):
    from switch_migrator import connection
    monkeypatch.setattr(connection.time, "sleep", lambda *_: None)

    conn = _FlakyConn(failures=1, output="PORT NUM\n1/1 up up")
    runner = _runner(conn, command_retries=1)
    assert runner.run("show interfaces gigabitEthernet state").startswith("PORT")
    assert conn.calls == 2
    assert runner.command_log[-1].attempts == 2 and runner.command_log[-1].ok
    # a recovered command must not count towards the circuit breaker
    assert runner._transport_failures == 0


def test_retries_are_bounded(monkeypatch):
    from switch_migrator import connection
    monkeypatch.setattr(connection.time, "sleep", lambda *_: None)

    conn = _FlakyConn(failures=99)
    runner = _runner(conn, command_retries=2)
    with pytest.raises(CommandError):
        runner.run("show mlt")
    assert conn.calls == 3                     # first try + two retries
    assert runner._transport_failures == 1     # one failed COMMAND, not three


def test_a_rejected_command_is_never_retried(monkeypatch):
    """'% Invalid input' is the device's final answer - re-sending it only
    wastes the read timeout on a box that already said no."""
    from switch_migrator import connection
    monkeypatch.setattr(connection.time, "sleep", lambda *_: None)

    conn = _FlakyConn(failures=0, output="% Invalid input detected at '^' marker.")
    runner = _runner(conn, command_retries=3)
    with pytest.raises(CommandError):
        runner.run("show interfaces gigabitEthernet state")
    assert conn.calls == 1


# --------------------------- progress hook ---------------------------------

def test_runner_reports_every_command_to_the_progress_hook(tmp_path, cfg):
    shutil.copytree(FIXTURES / "voss", tmp_path / "sw1")
    seen = []
    runner = OfflineRunner("sw1", tmp_path)
    runner.on_command = lambda device, command, index: seen.append(
        (device, command, index))
    collect_switch(SwitchTarget("sw1", "sw1", Platform.VOSS), runner, cfg)
    assert seen, "the collection sent commands but the hook never fired"
    assert all(d == "sw1" for d, _, _ in seen)
    assert [i for _, _, i in seen] == list(range(1, len(seen) + 1))
    assert len(runner.command_log) == len(seen)


def test_a_broken_progress_hook_never_breaks_the_collection(tmp_path, cfg):
    shutil.copytree(FIXTURES / "voss", tmp_path / "sw1")
    runner = OfflineRunner("sw1", tmp_path)

    def boom(*_a):
        raise RuntimeError("the progress bar exploded")

    runner.on_command = boom
    audit = collect_switch(SwitchTarget("sw1", "sw1", Platform.VOSS), runner, cfg)
    assert audit.ports and audit.reachable


def test_progress_is_silent_when_not_on_a_terminal():
    from rich.console import Console
    assert isinstance(make_progress(Console(force_terminal=False), 3),
                      NullProgress)


# --------------------------- manifest --------------------------------------

def test_manifest_records_commands_and_outcomes(tmp_path, cfg):
    shutil.copytree(FIXTURES / "voss", tmp_path / "sw1")
    runner = OfflineRunner("sw1", tmp_path)
    audit = collect_switch(SwitchTarget("sw1", "sw1", Platform.VOSS), runner, cfg)
    args = build_arg_parser().parse_args(["--no-fabric"])
    from datetime import datetime
    started = datetime(2026, 8, 12, 9, 0, 0)
    data = manifest_mod.build(
        [audit], FabricState(), args, started, datetime(2026, 8, 12, 9, 2, 30),
        [tmp_path / "report.xlsx"],
        {"sw1": [vars(r) for r in runner.command_log]},
        config_path=Path("config.yaml"), exit_code=0)

    assert data["duration_seconds"] == 150.0
    assert data["read_only"] is True
    device = data["switches"][0]
    assert device["name"] == "sw1" and device["reachable"]
    assert device["commands"]["sent"] == len(runner.command_log)
    assert device["counts"]["ports"] == len(audit.ports)
    assert data["totals"]["commands_sent"] == len(runner.command_log)
    assert data["files_written"] == [str(tmp_path / "report.xlsx")]


def test_manifest_never_carries_credentials(tmp_path, cfg):
    """The manifest is meant to be handed to other people, and it copies
    whatever the caller parsed. Credentials do not reach argparse today - but a
    file that WOULD print one if they ever did is a bad bet."""
    args = build_arg_parser().parse_args(["--no-fabric", "-s", "sw1:voss"])
    from datetime import datetime
    for name in ("password", "ssh_key", "api_token", "auth"):
        setattr(args, name, "topsecret")
    data = manifest_mod.build([], FabricState(), args, datetime.now(),
                              datetime.now(), [], {})
    assert "topsecret" not in json.dumps(data)
    assert data["options"]["password"] == "***redacted***"
    # and the ordinary options are still recorded
    assert data["options"]["switch"] == ["sw1:voss"]


# --------------------------- configurable knobs ----------------------------

def test_mac_cap_comes_from_the_config(tmp_path):
    path = tmp_path / "config.yaml"
    path.write_text("mac_cap: 3\nssh:\n  command_retries: 4\n"
                    "  max_transport_failures: 7\n")
    cfg = load_config(path, require_fabric=False)
    assert cfg.mac_cap == 3
    assert cfg.ssh.command_retries == 4
    assert cfg.ssh.max_transport_failures == 7


def test_mac_cap_defaults_and_is_sane(tmp_path):
    path = tmp_path / "config.yaml"
    path.write_text("mac_cap: 0\n")           # nonsense value must not disable it
    cfg = load_config(path, require_fabric=False)
    assert cfg.mac_cap == 1
    empty = tmp_path / "empty.yaml"
    empty.write_text("{}\n")
    assert load_config(empty, require_fabric=False).mac_cap == 10


def test_mac_cap_actually_caps_the_collected_macs(tmp_path, cfg):
    device = tmp_path / "sw1"
    shutil.copytree(FIXTURES / "voss", device)
    (device / "show_interfaces_gigabitethernet_fdb_entry.txt").write_text(
        "VLAN  MAC-ADDRESS        INTERFACE\n"
        + "".join(f"100   00:11:22:33:44:{n:02x} Port-1/1\n" for n in range(8)))
    cfg.mac_cap = 3
    audit = collect_switch(SwitchTarget("sw1", "sw1", Platform.VOSS),
                           OfflineRunner("sw1", tmp_path), cfg, pull_macs=True)
    port = next(p for p in audit.ports if p.port == "1/1")
    assert len(port.macs) == 3        # kept
    assert port.mac_total == 8        # but the true count is still reported


# --------------------------- dry-run runner --------------------------------

def test_dry_run_runner_answers_everything_with_nothing():
    runner = DryRunRunner("sw", Platform.VOSS)
    assert runner.run("show mlt") == ""
    assert runner.run("show vlan i-sid") == ""
    assert runner.commands == ["enable", "terminal more disable",
                               "term more dis", "show mlt", "show vlan i-sid"]
    assert all(r.ok for r in runner.command_log)


def test_progress_counts_devices_and_controllers():
    import io
    from rich.console import Console
    from switch_migrator.models import SwitchAudit
    from switch_migrator.report.progress import CollectionProgress

    console = Console(file=io.StringIO(), force_terminal=True, width=100)
    with CollectionProgress(console, switches=2, dvrs=1) as p:
        p.fabric_done("dvr-01", True)
        for name in ("sw-a", "sw-b"):
            p.device_start(name)
            p.device_command(name, "show mlt", 1)
        p.device_done(SwitchAudit(name="sw-a", host="a", platform=Platform.VOSS,
                                  reachable=True), 16)
        # an unreachable device still advances the bar, it just has nothing to
        # report - the CLI prints the UNREACHABLE line itself
        p.device_done(SwitchAudit(name="sw-b", host="b", platform=Platform.ERS,
                                  reachable=False), 0)
        tasks = {t.description: t for t in p._progress.tasks}
        overall = next(t for d, t in tasks.items() if "Auditing" in d)
        fabric = next(t for d, t in tasks.items() if "Fabric" in d)
        assert overall.completed == 2 and overall.total == 2
        assert fabric.completed == 1
        assert not p._tasks             # per-device rows are cleaned up


# --------------------------- location split (CLI) ---------------------------

def _location_env(tmp_path):
    """Three switches at two sites, replayed offline."""
    import shutil
    raw = tmp_path / "raw"
    raw.mkdir()
    for name in ("gx-11-s72-p1", "gx-12-s01-p9", "mu-01-a"):
        shutil.copytree(FIXTURES / "voss", raw / name)
    path = tmp_path / "config.yaml"
    path.write_text(
        "locations:\n"
        "  patterns:\n"
        "    'gx-11-*': Frankfurt DC1\n"
        "    'gx-12-*': Frankfurt DC2\n"
        "    'mu-*': Munich\n"
        "  groups:\n"
        "    Frankfurt: ['Frankfurt DC1', 'Frankfurt DC2']\n")
    switches = []
    for name in ("gx-11-s72-p1", "gx-12-s01-p9", "mu-01-a"):
        switches += ["-s", f"{name}:voss"]
    return path, raw, switches


def test_split_by_location_writes_one_sheet_per_group(tmp_path):
    cfg_path, raw, switches = _location_env(tmp_path)
    out = tmp_path / "out"
    main(["-c", str(cfg_path), "-o", str(out), "--offline", str(raw),
          "--no-fabric", "--migration-sheets", "--csv", "--no-excel",
          "--split-by-location"] + switches)
    csv_dir = next(out.glob("csv-*"))
    names = sorted(p.name for p in csv_dir.glob("cabling*.csv"))
    assert names == ["cabling_frankfurt.csv", "cabling_munich.csv"]
    # the combined sheet is gone: one link, one sheet to write it on
    assert not (csv_dir / "cabling.csv").exists()
    frankfurt = (csv_dir / "cabling_frankfurt.csv").read_text()
    assert "gx-11-s72-p1" in frankfurt and "gx-12-s01-p9" in frankfurt
    assert "mu-01-a" not in frankfurt


def test_location_group_overrides_the_config_for_one_run(tmp_path):
    cfg_path, raw, switches = _location_env(tmp_path)
    out = tmp_path / "out"
    main(["-c", str(cfg_path), "-o", str(out), "--offline", str(raw),
          "--no-fabric", "--migration-sheets", "--csv", "--no-excel",
          "--location-group", "Everything=Frankfurt DC*,Munich"] + switches)
    csv_dir = next(out.glob("csv-*"))
    assert sorted(p.name for p in csv_dir.glob("cabling*.csv")) == \
        ["cabling_everything.csv"]


def test_a_mistyped_location_group_stops_the_run(tmp_path):
    cfg_path, raw, switches = _location_env(tmp_path)
    code = main(["-c", str(cfg_path), "-o", str(tmp_path / "out"),
                 "--offline", str(raw), "--no-fabric", "--migration-sheets",
                 "--location-group", "Frankfurt gx-11"] + switches)
    assert code == 2            # a config error, not a silently split sheet


def test_without_the_flag_the_sheet_stays_combined(tmp_path):
    cfg_path, raw, switches = _location_env(tmp_path)
    out = tmp_path / "out"
    main(["-c", str(cfg_path), "-o", str(out), "--offline", str(raw),
          "--no-fabric", "--migration-sheets", "--csv", "--no-excel"] + switches)
    csv_dir = next(out.glob("csv-*"))
    assert (csv_dir / "cabling.csv").is_file()
    assert not list(csv_dir.glob("cabling_*.csv"))


def test_excel_tab_names_survive_long_and_illegal_location_names(tmp_path):
    from openpyxl import load_workbook
    from switch_migrator.report.excel import write_excel
    from switch_migrator.report.tables import Table

    tables = []
    for title in ("Cabling Frankfurt/Main DC1 - the whole east campus",
                  "Cabling Frankfurt/Main DC1 - the whole west campus",
                  "Cabling Munich [south]"):
        t = Table(title, ["Old switch", "Old port"])
        t.add(["sw", "1/1"])
        tables.append(t)
    path = tmp_path / "sheets.xlsx"
    write_excel(tables, path)
    from switch_migrator.report.excel import STAMP_SHEET
    # the hidden version stamp is not one of the report tabs
    titles = [t for t in load_workbook(path).sheetnames if t != STAMP_SHEET]
    assert len(titles) == 3, "a tab was lost to a name collision"
    assert len(set(titles)) == 3
    for title in titles:
        assert len(title) <= 31
        assert not set(title) & set("[]:*?/\\")


def test_cli_profile_supplies_inventory_and_settings(tmp_path, capsys):
    """--profile fills in what the command line left at defaults."""
    cfg_path = tmp_path / "config.yaml"
    cfg_path.write_text("excluded_vlans: [1]\n")
    inv = tmp_path / "inv.yaml"
    inv.write_text("switches:\n  - name: gx-01\n    platform: voss\n")
    profs = tmp_path / "profiles.yaml"
    profs.write_text(f"profiles:\n  site-a:\n    inventory: {inv}\n")
    code = main(["-c", str(cfg_path), "-o", str(tmp_path / "out"),
                 "--profiles-file", str(profs), "--profile", "site-a",
                 "--no-fabric", "--dry-run"])
    assert code == 0


def test_cli_unknown_profile_is_a_config_error(tmp_path):
    cfg_path = tmp_path / "config.yaml"
    cfg_path.write_text("excluded_vlans: [1]\n")
    code = main(["-c", str(cfg_path), "-o", str(tmp_path / "out"),
                 "--profiles-file", str(tmp_path / "profiles.yaml"),
                 "--profile", "nope", "--no-fabric", "--dry-run"])
    assert code == 2
