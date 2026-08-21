"""Reading a filled-in cabling sheet back in.

The sheet has been through a data centre and a spreadsheet program before it
comes back, so these tests are mostly about tolerance: added columns, moved
columns, hand-typed port ids, half-finished rows.
"""

import csv
from pathlib import Path

import pytest

from switch_migrator import cabling_sheet as CS

HEADERS = ["VLAN", "Type", "Port ID", "End device / neighbor", "NEW switch",
           "NEW port", "NEW MLT ID", "NEW MLT name", "NEW VLAN", "Old switch",
           "Old port", "MLT ID", "MLT name", "MLT VLANs", "MLT I-SIDs",
           "Port VLANs", "Port I-SIDs", "MAC addresses", "Media", "Usage", "Why"]


def _row(**kw) -> list:
    values = dict.fromkeys(HEADERS, "")
    values.update(kw)
    return [values[h] for h in HEADERS]


def _csv(tmp_path: Path, rows: list[list], headers: list[str] = None) -> Path:
    path = tmp_path / "cabling.csv"
    with path.open("w", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers or HEADERS)
        writer.writerows(rows)
    return path


def test_reads_the_columns_the_tool_wrote(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Port ID": "P0001", "Old switch": "gx-01", "Old port": "1/7",
                "NEW switch": "leaf-01", "NEW port": "1/7",
                "End device / neighbor": "srv-esx-01",
                "MAC addresses": "00:11:22:33:44:55,00:11:22:33:44:56",
                "Port VLANs": "695,735", "Port I-SIDs": "2500695,2510735",
                "MLT ID": "35", "MLT name": "MLT035", "Media": "10GbSR",
                "Usage": "IN USE", "Type": "mlt"}),
    ])
    sheet = CS.load(path)
    row = sheet.rows[0]
    assert row.uid == "P0001"
    assert (row.old_switch, row.old_port) == ("gx-01", "1/7")
    assert (row.new_switch, row.new_port) == ("leaf-01", "1/7")
    assert row.neighbor == "srv-esx-01"
    assert row.macs == ["00:11:22:33:44:55", "00:11:22:33:44:56"]
    assert row.port_vlans == [695, 735] and row.port_isids == [2500695, 2510735]
    assert row.mlt_id == 35 and row.mlt_name == "MLT035"
    assert row.media == "10GbSR" and row.usage == "IN USE" and row.kind == "mlt"
    assert row.migrated and sheet.problems == []


def test_extra_and_reordered_columns_are_fine(tmp_path):
    """People add their own columns and move things around. That must not
    break a file the whole migration depends on."""
    headers = ["Done by", "Old port", "NEW port", "Old switch", "NEW switch",
               "Rack", "Port ID"]
    path = _csv(tmp_path, [["TK", "1/7", "1/9", "gx-01", "leaf-01", "R14",
                            "P0001"]], headers)
    sheet = CS.load(path)
    row = sheet.rows[0]
    assert (row.old_switch, row.old_port) == ("gx-01", "1/7")
    assert (row.new_switch, row.new_port) == ("leaf-01", "1/9")
    assert row.uid == "P0001"
    assert sheet.problems           # says it kept columns it does not know


def test_hand_typed_port_ids_are_normalised(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": " Port 1/7 ",
                "NEW switch": "leaf-01", "NEW port": "1 / 9"}),
        _row(**{"Old switch": "ers-01", "Old port": "48",
                "NEW switch": "leaf-01", "NEW port": "GigabitEthernet 1/12"}),
    ])
    rows = CS.load(path).rows
    assert (rows[0].old_port, rows[0].new_port) == ("1/7", "1/9")
    assert (rows[1].old_port, rows[1].new_port) == ("48", "1/12")


def test_the_more_tail_of_a_mac_cell_is_not_read_as_a_mac(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": "1/7",
                "MAC addresses": "00:11:22:33:44:55 (+37 more)"})])
    assert CS.load(path).rows[0].macs == ["00:11:22:33:44:55"]


def test_unfilled_rows_are_pending_not_errors(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": "1/7"}),
        _row(**{"Old switch": "gx-01", "Old port": "1/8",
                "NEW switch": "leaf-01", "NEW port": "1/8"}),
    ])
    sheet = CS.load(path)
    assert len(sheet.pending) == 1 and len(sheet.migrated) == 1
    assert not sheet.pending[0].problems


def test_half_filled_rows_are_flagged(tmp_path):
    """A technician who started a line and got interrupted is the one case
    worth complaining about - it looks migrated but is not."""
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": "1/7",
                "NEW switch": "leaf-01"}),
        _row(**{"Old switch": "gx-01", "Old port": "1/8", "NEW port": "1/8"}),
    ])
    rows = CS.load(path).rows
    assert "NEW port is empty" in rows[0].problems[0]
    assert "NEW switch is empty" in rows[1].problems[0]
    assert not rows[0].migrated and not rows[1].migrated


def test_new_vlan_overrides_the_old_one(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": "1/7", "Port VLANs": "695",
                "NEW VLAN": "800"}),
        _row(**{"Old switch": "gx-01", "Old port": "1/8", "Port VLANs": "695"}),
        _row(**{"Old switch": "gx-01", "Old port": "1/9", "MLT VLANs": "200,300"}),
    ])
    rows = CS.load(path).rows
    assert rows[0].expected_vlans == [800]      # the planner's decision wins
    assert rows[1].expected_vlans == [695]
    assert rows[2].expected_vlans == [200, 300]  # MLT VLANs when the port has none


def test_a_nonsense_number_is_a_row_problem_not_a_crash(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": "1/7",
                "NEW switch": "leaf-01", "NEW port": "1/7",
                "NEW MLT ID": "ask Timo"})])
    row = CS.load(path).rows[0]
    assert row.new_mlt_id is None
    assert any("not a number" in p for p in row.problems)
    assert row.migrated             # the rest of the row is still usable


def test_blank_and_spacer_lines_are_skipped(tmp_path):
    path = _csv(tmp_path, [
        _row(**{"Old switch": "gx-01", "Old port": "1/7"}),
        [""] * len(HEADERS),
        _row(),
        _row(**{"Old switch": "gx-01", "Old port": "1/8"}),
    ])
    assert len(CS.load(path).rows) == 2


def test_a_file_that_is_not_a_cabling_sheet_says_so(tmp_path):
    path = tmp_path / "other.csv"
    path.write_text("name,value\nfoo,1\n")
    with pytest.raises(CS.SheetError, match="'Old switch', 'Old port'"):
        CS.load(path)


def test_clear_errors_for_missing_or_empty_files(tmp_path):
    with pytest.raises(CS.SheetError, match="not found"):
        CS.load(tmp_path / "nope.csv")
    empty = tmp_path / "empty.csv"
    empty.write_text("")
    with pytest.raises(CS.SheetError, match="empty"):
        CS.load(empty)
    headers_only = _csv(tmp_path, [])
    with pytest.raises(CS.SheetError, match="no data rows"):
        CS.load(headers_only)


# --------------------------- the real .xlsx --------------------------------

def test_reads_the_xlsx_the_tool_writes(tmp_path):
    """Round-trip through the actual writer: the sheet the technicians get is
    the sheet this has to read."""
    from switch_migrator.report.excel import write_excel
    from switch_migrator.report.tables import Table

    table = Table("Cabling", HEADERS)
    table.add(_row(**{"Port ID": "P0001", "Old switch": "gx-01",
                      "Old port": "1/7", "Port VLANs": "695",
                      "MAC addresses": "00:11:22:33:44:55"}))
    other = Table("Port Info", ["Port ID", "Switch"])
    other.add(["P0001", "gx-01"])
    path = tmp_path / "sheets.xlsx"
    write_excel([other, table], path)          # cabling is NOT the first sheet

    sheet = CS.load(path)
    assert len(sheet.rows) == 1
    assert sheet.rows[0].old_port == "1/7"
    assert sheet.rows[0].macs == ["00:11:22:33:44:55"]


def test_xlsx_sheet_can_be_named(tmp_path):
    from switch_migrator.report.excel import write_excel
    from switch_migrator.report.tables import Table

    table = Table("Patching", HEADERS)
    table.add(_row(**{"Old switch": "gx-01", "Old port": "1/7"}))
    path = tmp_path / "sheets.xlsx"
    write_excel([table], path)
    assert CS.load(path, sheet_name="Patching").rows[0].old_switch == "gx-01"
    with pytest.raises(CS.SheetError, match="no sheet named"):
        CS.load(path, sheet_name="Nope")


def test_every_cabling_worksheet_is_read_not_just_the_first(tmp_path):
    """--split-by-location writes one worksheet per site. Reading only the
    first would verify one site and silently ignore the rest - the migration
    would look complete while half of it was never checked."""
    from switch_migrator.report.excel import write_excel
    from switch_migrator.report.tables import Table

    tables = [Table("Summary", ["Switch", "Ports up"])]
    tables[0].add(["gx-11-s72-p1", "12/48"])          # not a cabling sheet
    for site, switch in (("Frankfurt", "gx-11-s72-p1"), ("Munich", "mu-01-a")):
        t = Table(f"Cabling {site}", HEADERS)
        t.add(_row(**{"Port ID": f"P-{site}", "Old switch": switch,
                      "Old port": "1/1", "NEW switch": "leaf-01",
                      "NEW port": "1/1"}))
        tables.append(t)
    path = tmp_path / "sheets.xlsx"
    write_excel(tables, path)

    sheet = CS.load(path)
    assert sheet.sheets_read == ["Cabling Frankfurt", "Cabling Munich"]
    assert {r.old_switch for r in sheet.rows} == {"gx-11-s72-p1", "mu-01-a"}
    # and each row remembers where it came from, for the error messages
    assert {r.sheet for r in sheet.rows} == {"Cabling Frankfurt",
                                             "Cabling Munich"}


def test_one_worksheet_can_still_be_singled_out(tmp_path):
    from switch_migrator.report.excel import write_excel
    from switch_migrator.report.tables import Table

    tables = []
    for site, switch in (("Frankfurt", "gx-11-s72-p1"), ("Munich", "mu-01-a")):
        t = Table(f"Cabling {site}", HEADERS)
        t.add(_row(**{"Old switch": switch, "Old port": "1/1"}))
        tables.append(t)
    path = tmp_path / "sheets.xlsx"
    write_excel(tables, path)
    sheet = CS.load(path, sheet_name="Cabling Munich")
    assert [r.old_switch for r in sheet.rows] == ["mu-01-a"]


def test_xlsx_worksheet_with_a_leading_blank_row_is_still_read(tmp_path):
    """openpyxl hands back leading blank rows; the header detection must skip
    them like the parser does, or the whole worksheet silently vanishes."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Cabling"
    ws.append([])                       # a stray empty first row
    ws.append(["Old switch", "Old port", "NEW switch", "NEW port"])
    ws.append(["gx-01", "1/7", "leaf-01", "1/1"])
    path = tmp_path / "sheet.xlsx"
    wb.save(path)
    sheet = CS.load(path)
    assert len(sheet.rows) == 1
    assert sheet.rows[0].old_port == "1/7"


# --------------------------------------------------------------------------- #
# The sheet the writer produces, read straight back in
# --------------------------------------------------------------------------- #

def _written_sheet(tmp_path, rows=None):
    from switch_migrator.models import UNTAGGED, PortState, SwitchAudit, VlanBinding
    from switch_migrator.models import Platform
    from switch_migrator.report.excel import write_excel
    from switch_migrator.report.migration import assign_port_uids, build_cabling

    a = SwitchAudit(name="old-01", host="old-01", platform=Platform.VOSS,
                    reachable=True)
    a.ports = [PortState(port="1/1", admin_up=True, oper_up=True,
                         lldp_neighbor="srv-a", usage="IN USE",
                         bindings=[VlanBinding(vlan=695, isid=2500695,
                                               tagging=UNTAGGED,
                                               source="running-config")])]
    assign_port_uids([a])
    table = build_cabling([a], ["new-01", "new-02"])
    path = tmp_path / "cabling.xlsx"
    write_excel([table], path)
    return path


def test_the_written_cabling_sheet_reads_back(tmp_path):
    path = _written_sheet(tmp_path)
    sheet = CS.load(path)
    assert len(sheet.rows) == 1
    row = sheet.rows[0]
    assert row.old_switch == "old-01" and row.old_port == "1/1"
    assert row.uid == "P0001"
    assert row.tagging == "untagged"
    # the rack columns exist and are empty - a human fills them in
    assert row.rack_old == "" and row.rack_new == ""
    # the hidden stamp is not mistaken for a worksheet of links
    assert "_migrator" not in sheet.sheets_read


def test_the_stamp_records_the_build_that_wrote_the_sheet(tmp_path):
    from switch_migrator.report.excel import SHEET_SCHEMA
    sheet = CS.load(_written_sheet(tmp_path))
    assert sheet.schema == SHEET_SCHEMA
    assert sheet.written_by
    # a current sheet raises no version complaint
    assert not [p for p in sheet.problems if "written by switch-migrator" in p]


def test_a_sheet_from_an_older_build_is_flagged_not_rejected(tmp_path):
    from openpyxl import load_workbook
    from switch_migrator.report.excel import STAMP_SHEET
    path = _written_sheet(tmp_path)
    wb = load_workbook(path)
    stamp = wb[STAMP_SHEET]
    for row in stamp.iter_rows():
        if row[0].value == "schema":
            row[1].value = "1"
    wb.save(path)

    sheet = CS.load(path)
    assert len(sheet.rows) == 1            # still read, not rejected
    assert any("sheet schema 1" in p for p in sheet.problems)


def test_rack_columns_are_read_when_filled_in(tmp_path):
    from openpyxl import load_workbook
    path = _written_sheet(tmp_path)
    wb = load_workbook(path)
    ws = wb["Cabling"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("Rack (old)") + 1).value = "R12"
    ws.cell(row=2, column=headers.index("Rack (new)") + 1).value = "R44"
    ws.cell(row=2, column=headers.index("NEW switch") + 1).value = "new-01"
    ws.cell(row=2, column=headers.index("NEW port") + 1).value = "Port 1/9"
    wb.save(path)

    row = CS.load(path).rows[0]
    assert row.rack_old == "R12" and row.rack_new == "R44"
    assert row.migrated and row.new_switch == "new-01"
    assert row.new_port == "1/9"           # normalised on the way in
